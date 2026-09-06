from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "PHASE-17-CEO-REVIEW-PACK.docx"
MAIN = ROOT / "PHASE-17-CEO-REVIEW-DRAFT.md"
APPENDIX = ROOT / "PHASE-17-EMAIL-TEMPLATE-REVIEW.md"
LOGO = ROOT / "public" / "full-logo.png"

NAVY = "0B2C6B"
PALE_BLUE = "F3F7FC"
PALE_GOLD = "FFF8E8"
GOLD = "D9A328"
MID_GRAY = "667085"
LIGHT_BORDER = "D9D9D9"
BLACK = "000000"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = LIGHT_BORDER, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        border = borders.find(qn(tag))
        if border is None:
            border = OxmlElement(tag)
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), size)
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, name="Aptos", size=None, bold=None, color=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), NAVY)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(color)
    run_properties.append(underline)
    run.append(run_properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\([^)]+\)|https?://\S+)")


def add_inline(paragraph, text: str, *, default_bold=False, default_color=BLACK) -> None:
    cursor = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, bold=default_bold, color=default_color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, bold=True, color=default_color)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Aptos Mono", size=9.5, color=NAVY)
        elif token.startswith("["):
            link = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if link and link.group(2).startswith("http"):
                add_hyperlink(paragraph, link.group(1), link.group(2))
            elif link:
                run = paragraph.add_run(link.group(1))
                set_run_font(run, bold=True, color=NAVY)
        else:
            url = token.rstrip(".,;)")
            suffix = token[len(url):]
            add_hyperlink(paragraph, url, url)
            if suffix:
                run = paragraph.add_run(suffix)
                set_run_font(run, color=default_color)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, bold=default_bold, color=default_color)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("1F2937")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    heading_sizes = {"Heading 1": 18, "Heading 2": 14, "Heading 3": 11.5}
    for name, size in heading_sizes.items():
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(14 if name != "Heading 1" else 20)
        style.paragraph_format.space_after = Pt(6)


def add_footer(doc: Document) -> None:
    for section in doc.sections:
        footer = section.footer
        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run("BinaHub   |   CEO Review Pack   |   ")
        set_run_font(run, size=8, color=MID_GRAY)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        paragraph._p.append(field)


def add_cover(doc: Document) -> None:
    if LOGO.exists():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(30)
        paragraph.add_run().add_picture(str(LOGO), width=Inches(2.15))

    label = doc.add_paragraph()
    label.paragraph_format.space_after = Pt(7)
    run = label.add_run("GOVERNANCE AND CONTROLLED PILOT")
    set_run_font(run, size=9, bold=True, color=GOLD)

    title = doc.add_paragraph(style="Title")
    title.add_run("Fase 17 Draft Review CEO")
    title_p_pr = title._p.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run("Keputusan bisnis yang perlu direview sebelum release dan controlled pilot")
    set_run_font(run, size=14, color=MID_GRAY)

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(18)
    add_inline(
        intro,
        "Dokumen ini sudah berisi rekomendasi awal. CEO cukup memilih SETUJUI, REVISI, atau TUNDA dan menuliskan koreksi yang diperlukan. Persetujuan tidak mengaktifkan automation atau outbound.",
    )

    metadata = doc.add_table(rows=4, cols=2)
    metadata.alignment = WD_TABLE_ALIGNMENT.LEFT
    metadata.autofit = False
    metadata.columns[0].width = Inches(1.65)
    metadata.columns[1].width = Inches(4.85)
    values = [
        ("Ditujukan kepada", "CEO BinaHub"),
        ("Disiapkan", "5 September 2026"),
        ("Status", "Draft untuk review"),
        ("Safety", "Automation tetap dry-run dan inactive"),
    ]
    for index, (label_text, value_text) in enumerate(values):
        row = metadata.rows[index]
        for cell in row.cells:
            set_cell_borders(cell)
            set_cell_margins(cell, top=130, bottom=130)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(row.cells[0], PALE_BLUE)
        left = row.cells[0].paragraphs[0]
        right = row.cells[1].paragraphs[0]
        add_inline(left, label_text, default_bold=True, default_color=NAVY)
        add_inline(right, value_text)

    doc.add_paragraph()
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(18)
    note.paragraph_format.space_after = Pt(0)
    run = note.add_run("Dokumen terdiri dari lembar keputusan utama dan lampiran 18 template email.")
    set_run_font(run, size=9.5, color=MID_GRAY)
    doc.add_page_break()


def parse_table(lines: list[str], start: int):
    rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        rows.append(cells)
        index += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", value or "") for value in rows[1]):
        rows.pop(1)
    return rows, index


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    repeat_table_header(table.rows[0])

    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        keep_row_together(row)
        for column_index in range(column_count):
            cell = row.cells[column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            set_cell_margins(cell)
            value = values[column_index] if column_index < len(values) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            if row_index == 0:
                set_cell_shading(cell, NAVY)
                add_inline(paragraph, value, default_bold=True, default_color="FFFFFF")
            else:
                if row_index % 2 == 0:
                    set_cell_shading(cell, PALE_BLUE)
                add_inline(paragraph, value)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_markdown(doc: Document, path: Path, *, appendix=False) -> None:
    lines = path.read_text(encoding="utf-8").splitlines()
    index = 0
    if appendix:
        doc.add_page_break()
    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("# "):
            if appendix:
                paragraph = doc.add_paragraph(style="Heading 1")
                paragraph.paragraph_format.page_break_before = True
                add_inline(paragraph, "Lampiran Review 18 Template Email", default_bold=True)
            index += 1
            continue
        if stripped.startswith("## "):
            heading = re.sub(r"^##\s+", "", stripped).replace("—", " ").replace("&", "dan")
            paragraph = doc.add_paragraph(style="Heading 1")
            if appendix and heading.startswith(("A. ", "B. ", "C. ")):
                paragraph.paragraph_format.page_break_before = True
            if appendix and heading == "Persetujuan akhir lampiran":
                paragraph.paragraph_format.page_break_before = True
            add_inline(paragraph, heading)
            index += 1
            continue
        if stripped.startswith("### "):
            heading = re.sub(r"^###\s+", "", stripped).replace("—", " ")
            paragraph = doc.add_paragraph(style="Heading 2")
            add_inline(paragraph, heading)
            index += 1
            continue
        if stripped.startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(doc, rows)
            continue
        if stripped.startswith("> "):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.right_indent = Inches(0.15)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(8)
            add_inline(paragraph, stripped[2:], default_color=NAVY)
            index += 1
            continue
        if re.match(r"^[-*]\s+", stripped):
            item = re.sub(r"^[-*]\s+", "", stripped)
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.left_indent = Inches(0.22)
            paragraph.paragraph_format.first_line_indent = Inches(-0.14)
            if item.startswith("[ ] "):
                item = "☐ " + item[4:]
                paragraph.style = doc.styles["Normal"]
                paragraph.paragraph_format.left_indent = Inches(0.12)
                paragraph.paragraph_format.space_after = Pt(7)
            add_inline(paragraph, item)
            index += 1
            continue
        numbered = re.match(r"^(\d+)\.\s+(.+)", stripped)
        if numbered:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.24)
            paragraph.paragraph_format.first_line_indent = Inches(-0.16)
            add_inline(paragraph, f"{numbered.group(1)}. {numbered.group(2)}")
            index += 1
            continue
        if stripped.startswith("Tanggal disiapkan:") or stripped.startswith("Status:") or stripped.startswith("Versi sumber:"):
            index += 1
            continue
        paragraph = doc.add_paragraph()
        if stripped.startswith("**Subject:**"):
            paragraph.paragraph_format.keep_with_next = True
        add_inline(paragraph, stripped)
        index += 1


def main() -> None:
    document = Document()
    style_document(document)
    add_cover(document)
    add_markdown(document, MAIN)
    add_markdown(document, APPENDIX, appendix=True)
    add_footer(document)

    document.core_properties.title = "Fase 17 Draft Review CEO"
    document.core_properties.subject = "Governance and controlled pilot review"
    document.core_properties.author = "BinaHub"
    document.core_properties.keywords = "BinaHub, CEO, governance, pilot, review"
    document.core_properties.comments = "Prepared for CEO review"

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
